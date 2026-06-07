"""Загрузчики PDF/DOCX/TXT. Используют pypdf+pdfplumber и python-docx."""

from __future__ import annotations

from pathlib import Path

from pydantic import BaseModel

from methodist_chat.infra.logging import get_logger
from .base import Tool, ToolResult
from .registry import register

logger = get_logger(__name__)


class LoadRequest(BaseModel):
    path: str


class LoadResponse(BaseModel):
    text: str
    char_count: int
    pages: int = 0
    fmt: str = ""


def _load_pdf(path: Path) -> tuple[str, int]:
    """Сначала pypdf, при пустом тексте — fallback на pdfplumber."""
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        chunks: list[str] = []
        for page in reader.pages:
            try:
                chunks.append(page.extract_text() or "")
            except Exception:  # pragma: no cover
                chunks.append("")
        text = "\n".join(chunks).strip()
        pages = len(reader.pages)
    except Exception as e:  # pragma: no cover
        logger.warning(f"pypdf failed: {e}; fallback to pdfplumber")
        text, pages = "", 0

    if len(text) < 50:
        try:
            import pdfplumber

            with pdfplumber.open(str(path)) as pdf:
                pages = len(pdf.pages)
                text = "\n".join((p.extract_text() or "") for p in pdf.pages).strip()
        except Exception as e:  # pragma: no cover
            logger.warning(f"pdfplumber failed: {e}")
    return text, pages


def _load_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts: list[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    # тексты из таблиц тоже бывают полезны
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                t = cell.text.strip()
                if t:
                    parts.append(t)
    return "\n".join(parts)


def load_document(path: str | Path) -> LoadResponse:
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(p)
    suffix = p.suffix.lower()
    if suffix == ".pdf":
        text, pages = _load_pdf(p)
        return LoadResponse(text=text, char_count=len(text), pages=pages, fmt="pdf")
    if suffix in {".docx", ".doc"}:
        text = _load_docx(p)
        return LoadResponse(text=text, char_count=len(text), fmt="docx")
    if suffix in {".txt", ".md"}:
        text = p.read_text(encoding="utf-8")
        return LoadResponse(text=text, char_count=len(text), fmt="txt")
    raise ValueError(f"Неподдерживаемый формат: {suffix}")


@register("doc_loader")
class DocLoaderTool(Tool[LoadRequest, LoadResponse]):
    description = "Загружает PDF/DOCX/TXT с диска и возвращает извлечённый текст."

    def run(self, params: LoadRequest) -> ToolResult[LoadResponse]:
        data = load_document(params.path)
        return ToolResult(
            data=data, log=[f"loaded {data.fmt} chars={data.char_count} pages={data.pages}"]
        )
